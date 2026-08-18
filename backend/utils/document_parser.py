from __future__ import annotations

import io
import re
from pathlib import Path

import docx
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class DocumentParsingError(ValueError):
    """Raised when an uploaded document cannot be parsed."""


def extract_text(uploaded_file) -> str:
    """Extract text from an uploaded PDF, DOCX, or TXT file."""
    file_name = getattr(uploaded_file, "filename", None) or getattr(uploaded_file, "name", "") or ""
    extension = Path(file_name).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentParsingError(
            "Unsupported document type. Please upload a PDF, DOCX, or TXT file."
        )

    try:
        payload = uploaded_file.getvalue() if hasattr(uploaded_file, "getvalue") else uploaded_file.read()
        if extension == ".pdf":
            return _extract_pdf(payload)
        if extension == ".docx":
            return _extract_docx(payload)
        return _extract_txt(payload)
    except DocumentParsingError:
        raise
    except Exception as exc:
        raise DocumentParsingError(f"Could not read the uploaded document: {exc}") from exc


def clean_text(text: str) -> str:
    """Normalize extracted legal text while preserving paragraph boundaries."""
    if not text:
        return ""

    text = text.replace("\x00", " ")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _extract_pdf(payload: bytes) -> str:
    reader = PdfReader(io.BytesIO(payload))
    if reader.is_encrypted:
        raise DocumentParsingError(
            "The PDF is encrypted. Please upload an unlocked document."
        )

    pages = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        if page_text.strip():
            pages.append(page_text)

    text = "\n\n".join(pages).strip()
    if not text:
        raise DocumentParsingError(
            "No selectable text was found in this PDF. It may be scanned; OCR is required before analysis."
        )
    return text


def _extract_docx(payload: bytes) -> str:
    document = docx.Document(io.BytesIO(payload))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))

    text = "\n\n".join(paragraphs).strip()
    if not text:
        raise DocumentParsingError("The DOCX file did not contain readable text.")
    return text


def _extract_txt(payload: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            text = payload.decode(encoding)
            if text.strip():
                return text
        except UnicodeDecodeError:
            continue
    raise DocumentParsingError("The TXT file is empty or uses an unsupported encoding.")
